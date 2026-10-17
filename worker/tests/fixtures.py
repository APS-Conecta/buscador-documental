# SPDX-License-Identifier: AGPL-3.0-or-later
"""Generates test documents on the fly. Hostile files are never committed."""
import io
import zipfile

import pikepdf
from openpyxl import Workbook
from PIL import Image

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def txt(path, text, encoding="utf-8"):
    path.write_bytes(text.encode(encoding))
    return str(path)


def _zip(path, files, compress=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(path, "w", compress) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return str(path)


def docx(path, paragraphs, table=None):
    body = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphs)
    if table:
        rows = "".join("<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in r) + "</w:tr>" for r in table)
        body += f"<w:tbl>{rows}</w:tbl>"
    xml = f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'
    return _zip(path, {"[Content_Types].xml": "<Types/>", "word/document.xml": xml})


def pptx(path, slides):
    files = {"[Content_Types].xml": "<Types/>"}
    for i, texts in enumerate(slides, start=1):
        runs = "".join(f"<a:p><a:r><a:t>{t}</a:t></a:r></a:p>" for t in texts)
        files[f"ppt/slides/slide{i}.xml"] = f'<?xml version="1.0"?><p:sld xmlns:p="p" xmlns:a="{A}"><p:txBody>{runs}</p:txBody></p:sld>'
    return _zip(path, files)


def odt(path, paragraphs):
    body = "".join(f"<text:p>{p}</text:p>" for p in paragraphs)
    xml = f'<?xml version="1.0"?><office:document-content xmlns:office="o" xmlns:text="t"><office:body><office:text>{body}</office:text></office:body></office:document-content>'
    return _zip(path, {"mimetype": "application/vnd.oasis.opendocument.text", "content.xml": xml})


def xlsx(path, sheets):
    wb = Workbook(write_only=True)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return str(path)


def pdf_text(path, pages=1, text="Hola mundo", encrypt=None):
    pdf = pikepdf.new()
    for i in range(pages):
        pdf.add_blank_page(page_size=(200, 200))
        page = pdf.pages[-1]
        content = f"BT /F1 12 Tf 20 100 Td ({text} {i + 1}) Tj ET".encode()
        page.Contents = pdf.make_stream(content)
        page.Resources = pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=pikepdf.Dictionary(Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1, BaseFont=pikepdf.Name.Helvetica)))
    kw = {}
    if encrypt:
        kw["encryption"] = pikepdf.Encryption(owner=encrypt[0], user=encrypt[1])
    pdf.save(path, **kw)
    return str(path)


# ---- hostile ----
def zip_bomb_docx(path):
    big = b"<w:t>" + b"a" * (200 * 1024 * 1024) + b"</w:t>"        # compresses ~1000:1
    return _zip(path, {"[Content_Types].xml": "<Types/>", "word/document.xml": big})


def nested_xml_docx(path, depth=5000):
    xml = '<w:document xmlns:w="%s">' % W + "<w:p>" * depth + "<w:t>x</w:t>" + "</w:p>" * depth + "</w:document>"
    return _zip(path, {"[Content_Types].xml": "<Types/>", "word/document.xml": xml})


def entity_bomb_docx(path):
    xml = ('<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">'
           '<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">]><w:document xmlns:w="%s"><w:body><w:p><w:r><w:t>&c;</w:t></w:r></w:p></w:body></w:document>' % W)
    return _zip(path, {"[Content_Types].xml": "<Types/>", "word/document.xml": xml})


def truncated_zip(path):
    good = io.BytesIO()
    with zipfile.ZipFile(good, "w") as z:
        z.writestr("word/document.xml", "<x/>" * 1000)
    path.write_bytes(good.getvalue()[:-30])
    return str(path)


def pixel_bomb_png(path):
    Image.new("1", (10000, 10000)).save(path, "PNG", optimize=True)    # 100 MP, a few KB
    return str(path)


def pixel_bomb_pdf(path):
    pdf = pikepdf.new()
    pdf.add_blank_page(page_size=(200, 200))
    img = pikepdf.Stream(pdf, b"\x00", Type=pikepdf.Name.XObject, Subtype=pikepdf.Name.Image, Width=20000, Height=20000,
                         ColorSpace=pikepdf.Name.DeviceGray, BitsPerComponent=1)
    pdf.pages[0].Resources = pikepdf.Dictionary(XObject=pikepdf.Dictionary(Im0=img))
    pdf.pages[0].Contents = pdf.make_stream(b"q 200 0 0 200 0 0 cm /Im0 Do Q")
    pdf.save(path)
    return str(path)


def big_xlsx(path, cells=1_200_000):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("s")
    for _ in range(cells // 100):
        ws.append(list(range(100)))
    wb.save(path)
    return str(path)


# ---- scanned pages ----
LINES = ["Informe de laboratorio", "", "Hemoglobina glicosilada 6,8 por ciento.",
         "La paciente fue derivada a nutricion y control en tres meses."]


def scan_image(lines=LINES, rotate=0):
    from PIL import ImageDraw, ImageFont
    img = Image.new("L", (1654, 2339), 255)
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 44)
    y = 200
    for line in lines:
        d.text((150, y), line, font=font, fill=0)
        y += 90
    return img.rotate(rotate, expand=True) if rotate else img


def scan_pdf(path, lines=LINES, rotate=0, garbage_layer=False, layer_text=None, visible=False):
    """An image-only page, optionally rotated, optionally with a text layer on top (garbage, or `layer_text`)."""
    garbage_layer = garbage_layer or layer_text is not None
    if not garbage_layer:
        scan_image(lines, rotate).save(path, "PDF", resolution=200.0)
        return str(path)
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    img = scan_image(lines, rotate)
    w, h = img.size
    c = canvas.Canvas(str(path), pagesize=(w * 72 / 200, h * 72 / 200))
    c.drawImage(ImageReader(img), 0, 0, width=w * 72 / 200, height=h * 72 / 200)
    t = c.beginText(40, 400)
    t.setTextRenderMode(0 if visible else 3)     # 3 = invisible
    t.textLine(layer_text if layer_text is not None else "#@!$%^&*()_+{}|:<>?~" * 6)
    c.drawText(t)
    c.save()
    return str(path)


def merge_pdfs(path, parts):
    out = pikepdf.new()
    for p in parts:
        with pikepdf.open(p) as src:
            out.pages.extend(src.pages)
    out.save(path)
    return str(path)


def native_with_logo(path, text="La paciente fue derivada a nutricion y control en tres meses"):
    """A born-digital page that carries a small logo image: it must stay a text-layer page."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    c.drawString(60, 700, text)
    c.drawImage(ImageReader(Image.new("RGB", (120, 60), (200, 30, 30))), 60, 760, width=120, height=60)
    c.save()
    return str(path)
